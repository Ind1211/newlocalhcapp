"""
Extracts plain text from an uploaded document so it can be included directly in an
AI prompt — no vector store, no chunking. A typical patient's document set is
small enough to fit in a single prompt alongside an assessment's field schema,
which is why this deliberately stays this simple rather than building toward RAG.
"""

import io

from docx import Document
from pypdf import PdfReader


class ExtractionError(Exception):
    pass


def extract_text(content: bytes, content_type: str, filename: str) -> str:
    if content_type == "application/pdf" or filename.lower().endswith(".pdf"):
        return _extract_pdf(content)
    if content_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ) or filename.lower().endswith(".docx"):
        return _extract_docx(content)
    if content_type.startswith("text/") or filename.lower().endswith(".txt"):
        return _extract_plain_text(content)
    raise ExtractionError(
        f"Unsupported file type '{content_type}' — supported: PDF, DOCX, and plain text."
    )


def _extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(p for p in pages if p.strip())
    except Exception as e:
        raise ExtractionError(f"Could not read this PDF: {e}") from e
    if not text.strip():
        raise ExtractionError(
            "No extractable text found — this PDF may be a scanned image rather than "
            "real text. Scanned-document support isn't built yet."
        )
    return text


def _extract_docx(content: bytes) -> str:
    try:
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Tables are common in clinical documents (lab result grids, medication
        # lists) — skipping them would silently drop real content.
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        text = "\n".join(paragraphs)
    except Exception as e:
        raise ExtractionError(f"Could not read this document: {e}") from e
    if not text.strip():
        raise ExtractionError("This document appears to be empty.")
    return text


def _extract_plain_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as e:
        raise ExtractionError(f"Could not decode this file as text: {e}") from e
