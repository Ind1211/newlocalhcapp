"""
Document upload and retry-extraction — object storage is mocked at the
storage-function level (app.storage.upload_document/download_document), the
same approach used throughout this project's own manual verification, rather
than standing up a real Azure Blob Storage / Azurite instance for tests to
depend on.
"""

from io import BytesIO
from unittest.mock import patch

from app.models import Patient, PatientDocument


def _fake_docx_bytes() -> bytes:
    """A minimal, real .docx — python-docx can create one in memory, which is
    what document_extraction.py actually needs to successfully extract text
    from, rather than an arbitrary byte string that would just fail differently."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("Patient has known chronic kidney disease.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_retry_extraction_succeeds_on_a_previously_failed_document(client, staff_factory, auth_headers, db):
    """Simulates the real, useful case for this endpoint: a document that
    failed extraction the first time (e.g. a transient storage hiccup) can be
    retried without re-uploading — and this time succeeds."""
    staff, token = staff_factory("nurse")
    headers = auth_headers(token)

    patient = Patient(name="Retry Test Patient", mrn="RETRY-001")
    db.add(patient)
    db.commit()

    fake_bytes = _fake_docx_bytes()
    doc = PatientDocument(
        patient_id=patient.id, filename="referral.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_key="fake-key-123", size_bytes=len(fake_bytes),
        extraction_status="failed", extraction_error="Simulated earlier failure",
    )
    db.add(doc)
    db.commit()

    with patch("app.routers.documents.download_document", return_value=fake_bytes):
        resp = client.post(f"/api/patients/{patient.id}/documents/{doc.id}/retry-extraction", headers=headers)

    assert resp.status_code == 200
    body = resp.json()
    assert body["extraction_status"] == "ok"
    assert "chronic kidney disease" in body["extracted_text"]


def test_retry_extraction_on_nonexistent_document_returns_404(client, staff_factory, auth_headers, db):
    staff, token = staff_factory("nurse")
    headers = auth_headers(token)
    patient = Patient(name="No Doc Patient", mrn="RETRY-002")
    db.add(patient)
    db.commit()

    resp = client.post(f"/api/patients/{patient.id}/documents/00000000-0000-0000-0000-000000000000/retry-extraction", headers=headers)
    assert resp.status_code == 404
